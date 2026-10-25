import json
import tempfile
import unittest
from pathlib import Path

from workshop.cards import (card, card_label, from_pbn, hcp, parse_card,
                            parse_suit_group, shape_text, sorted_cards, to_pbn)
from workshop.commands import CommandError, run_command
from workshop.model import Contract, Deal, DefensiveOverride
from workshop.storage import (FileFormatError, append_to_library,
                              append_to_library_text, check_out_of_library,
                              contract_text, deal_from_dict, from_tutor_dict,
                              library_deals, library_names, load_any,
                              LibraryText, read_library,
                              remove_from_library_text, replace_library,
                              save_workshop, to_library_dict, to_tutor_dict,
                              tutor_json_text)

SAMPLE = {"N": "AQ72.K63.984.A52", "E": "865.42.765.KT987",
          "S": "KT3.AQJT9.A32.Q4", "W": "J94.875.KQJT.J63"}


def sample_deal():
    d = Deal()
    for seat, pbn in SAMPLE.items():
        d.set_hand(seat, from_pbn(pbn))
    d.contract = Contract(level=4, strain="H", declarer="S")
    d.opening_lead = parse_card("DK")
    d.history.clear()
    return d


class CardTests(unittest.TestCase):
    def test_numbering_matches_tutor(self):
        self.assertEqual(parse_card("SA"), 12)
        self.assertEqual(parse_card("H2"), 13)
        self.assertEqual(parse_card("KD"), 37)
        self.assertEqual(parse_card("♣A"), 51)
        self.assertEqual(parse_card("10D"), card(2, 8))

    def test_suit_group(self):
        self.assertEqual(parse_suit_group("SAKQ"), [12, 11, 10])
        self.assertEqual(parse_suit_group("D10"), [card(2, 8)])

    def test_pbn_round_trip(self):
        cards = from_pbn("AKQ7.J94.K5.432")
        self.assertEqual(to_pbn(cards), "AKQ7.J94.K5.432")
        self.assertEqual(hcp(cards), 9 + 1 + 3)
        self.assertEqual(shape_text(cards), "4-3-2-3")
        self.assertEqual(to_pbn(from_pbn("A10 . . . ")), "AT...")
        with self.assertRaises(ValueError):
            from_pbn("AA...")


class DealTests(unittest.TestCase):
    def test_complete_deal_validates(self):
        v = sample_deal().validation()
        self.assertTrue(v.complete)
        self.assertEqual(v.errors, [])
        self.assertEqual(v.warnings, [])

    def test_swap_and_undo(self):
        d = sample_deal()
        q, seven = parse_card("DQ"), parse_card("D7")   # West, East
        self.assertTrue(d.swap(q, seven))
        self.assertEqual(d.location[q], "E")
        self.assertEqual(d.location[seven], "W")
        self.assertFalse(d.swap(q, parse_card("D6")))  # both East
        d.undo()
        self.assertEqual(d.location[q], "W")
        d.redo()
        self.assertEqual(d.location[q], "E")

    def test_move_counts_and_lead_warning(self):
        d = sample_deal()
        d.move([parse_card("DK")], "E")
        v = d.validation()
        self.assertFalse(v.complete)
        self.assertIn("West has 12 cards", v.errors)
        self.assertTrue(any("Opening lead" in w for w in v.warnings))

    def test_set_hand_pulls_cards_and_releases_others(self):
        d = sample_deal()
        d.set_hand("N", from_pbn("AQ72.K63.984.A5Q"))   # take ♣Q from South, drop ♣2
        self.assertEqual(d.location[parse_card("CQ")], "N")
        self.assertIsNone(d.location[parse_card("C2")])
        self.assertEqual(len(d.hand("S")), 12)

    def test_clear_deal_empties_hands_and_leads_and_undoes(self):
        d = sample_deal()
        d.later_leads["W"] = [parse_card("DQ")]
        d.add_override(DefensiveOverride("C", 1, "E", parse_card("CK")))
        d.author.title = "Hand 12"
        self.assertTrue(d.clear_deal())
        self.assertEqual(d.pool(), list(sorted_cards(range(52))))
        self.assertTrue(all(d.hand(s) == [] for s in "NESW"))
        self.assertIsNone(d.opening_lead)
        self.assertEqual(d.later_leads, {"W": [], "E": []})
        self.assertEqual(d.contract, Contract(4, "H", "S"))          # kept
        self.assertEqual(d.author.title, "Hand 12")                  # kept
        self.assertEqual(len(d.defensive_overrides), 1)              # kept
        self.assertFalse(d.clear_deal())                             # nothing left to clear
        d.undo()
        self.assertEqual(to_pbn(d.hand("N")), SAMPLE["N"])
        self.assertEqual(d.opening_lead, parse_card("DK"))
        self.assertEqual(d.later_leads["W"], [parse_card("DQ")])

    def test_versions(self):
        d = sample_deal()
        d.snapshot_version("base")
        d.clear_all()
        self.assertEqual(len(d.pool()), 52)
        d.restore_version(0)
        self.assertEqual(to_pbn(d.hand("N")), SAMPLE["N"])


class StorageTests(unittest.TestCase):
    def test_workshop_round_trip(self):
        d = sample_deal()
        d.author.title = "Hand 12"
        d.author.notes = "Ducking ♦"
        d.later_leads["W"] = [parse_card("DQ")]
        d.snapshot_version("base")
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "x.json"
            save_workshop(d, p)
            d2, kind = load_any(p)
        self.assertEqual(kind, "workshop")
        self.assertEqual(d2.hands(), d.hands())
        self.assertEqual(d2.contract, d.contract)
        self.assertEqual(d2.opening_lead, d.opening_lead)
        self.assertEqual(d2.later_leads, d.later_leads)
        self.assertEqual(d2.author, d.author)
        self.assertEqual(d2.versions[0].hands, d.versions[0].hands)

    def test_workshop_file_holds_card_numbers_only(self):
        d = sample_deal()
        d.move([parse_card("C2")], None)              # one card left in the pool
        d.snapshot_version("base")
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "x.json"
            save_workshop(d, p)
            data = json.loads(p.read_text(encoding="utf-8"))
            self.assertEqual(data["version"], 2)
            self.assertEqual(data["hands"]["N"], [12, 10, 5, 0, 24, 17, 14, 33, 32, 28, 51, 42])
            self.assertEqual(data["pool"], [39])
            self.assertEqual(data["versions"][0]["hands"]["N"], data["hands"]["N"])
            for hand in list(data["hands"].values()) + list(data["versions"][0]["hands"].values()):
                self.assertTrue(all(isinstance(c, int) for c in hand))
            d2, _ = load_any(p)
            self.assertEqual(d2.hands(), d.hands())
            self.assertEqual(d2.versions[0].hands, d.versions[0].hands)
            # Version 1 files stored hands as PBN text; they still load the same deal.
            data.update(version=1, hands=SAMPLE, pool="")
            p.write_text(json.dumps(data), encoding="utf-8")
            old, _ = load_any(p)
            self.assertEqual(old.hands(), sample_deal().hands())
            data["hands"] = {**SAMPLE, "N": [12, 99]}
            p.write_text(json.dumps(data), encoding="utf-8")
            with self.assertRaises(FileFormatError):
                load_any(p)

    def test_workshop_leads_saved_as_numbers_and_old_codes_load(self):
        d = sample_deal()
        d.opening_lead = 0                          # the two of spades: falsy but valid
        d.later_leads["W"] = [parse_card("DQ")]
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "x.json"
            save_workshop(d, p)
            data = json.loads(p.read_text(encoding="utf-8"))
            self.assertEqual(data["opening_lead"], 0)
            self.assertEqual(data["later_leads"], {"W": [36], "E": []})
            self.assertEqual(load_any(p)[0].opening_lead, 0)
            data.update(opening_lead="DK", later_leads={"W": ["DQ"], "E": []})   # older file
            p.write_text(json.dumps(data), encoding="utf-8")
            old, _ = load_any(p)
        self.assertEqual(old.opening_lead, parse_card("DK"))
        self.assertEqual(old.later_leads["W"], [parse_card("DQ")])

    def test_tutor_export_fields(self):
        d = sample_deal()
        d.later_leads["W"] = [parse_card("DQ")]
        out = to_tutor_dict(d)
        self.assertEqual(out["contract"], "4 Hearts")
        self.assertEqual(out["contract_trump"], "H")
        self.assertEqual(out["w_leads"], [parse_card("DK"), parse_card("DQ")])
        self.assertEqual(out["west_leads"], out["w_leads"])
        self.assertEqual(out["si_player"], 4)
        self.assertEqual(out["south_cards"][0], parse_card("SK"))   # spades first, high first
        json.loads(tutor_json_text(d))    # valid JSON

    def test_contract_text(self):
        self.assertEqual(contract_text(Contract(1, "S")), "1 Spade")
        self.assertEqual(contract_text(Contract(3, "N")), "3 NoTrumps")

    def test_tutor_import_with_duplicates(self):
        data = {"south_cards": [12, 11], "west_cards": [12, 9], "north_cards": [],
                "east_cards": [], "contract": "5 Clubs", "west_leads": [9],
                "east_leads": [], "si_action": 2, "commentary": "hi"}
        d = from_tutor_dict(data)
        self.assertEqual(d.duplicates, [12])
        self.assertEqual(d.location[12], "S")
        self.assertEqual(d.contract, Contract(5, "C", "S"))
        self.assertEqual(d.opening_lead, 9)
        out = to_tutor_dict(d)
        self.assertEqual(out["si_action"], 2)
        self.assertEqual(out["commentary"], "hi")

    def test_real_tutor_files_round_trip(self):
        """Import every tutor file found next to this project and re-export it."""
        root = Path(__file__).resolve().parents[2]
        files = [root / "BWARepo" / "1a Mendelson7.json",
                 root / "Bridge_Declarer Web App" / "Mendelson33",
                 root / "Book" / "Hands" / "Hand45.json"]
        files = [f for f in files if f.exists()]
        if not files:
            self.skipTest("no sample tutor files on this machine")
        for f in files:
            with self.subTest(file=f.name):
                original = json.loads(f.read_text(encoding="utf-8-sig"))
                d, kind = load_any(f)
                self.assertEqual(kind, "tutor")
                out = to_tutor_dict(d)
                for key in ("south_cards", "west_cards", "north_cards", "east_cards"):
                    self.assertEqual(sorted(out[key]), sorted(original[key]))
                orig_w = original.get("w_leads") or original.get("west_leads") or []
                self.assertEqual(out["w_leads"], orig_w)


class OverrideTests(unittest.TestCase):
    def deal_with_overrides(self):
        d = sample_deal()
        d.add_override(DefensiveOverride("D", 1, "W", parse_card("DQ")))
        d.add_override(DefensiveOverride("C", 2, "E", parse_card("CK")))
        return d

    EXPECTED = [{"player": "West", "suitLed": "D", "occurrence": 1, "cardToPlay": 36},
                {"player": "East", "suitLed": "C", "occurrence": 2, "cardToPlay": 50}]

    def test_saved_as_card_numbers_in_every_format(self):
        d = self.deal_with_overrides()
        self.assertEqual(to_tutor_dict(d)["defensiveOverrides"], self.EXPECTED)
        self.assertEqual(to_library_dict(d)["defensiveOverrides"], self.EXPECTED)
        self.assertEqual(json.loads(tutor_json_text(d))["defensiveOverrides"], self.EXPECTED)

    def test_workshop_round_trip(self):
        d = self.deal_with_overrides()
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "x.json"
            save_workshop(d, p)
            self.assertEqual(json.loads(p.read_text(encoding="utf-8"))["defensiveOverrides"],
                             self.EXPECTED)
            d2, _ = load_any(p)
        self.assertEqual(d2.defensive_overrides, d.defensive_overrides)

    def test_tutor_import_keeps_overrides_out_of_extra(self):
        d2 = from_tutor_dict(to_tutor_dict(self.deal_with_overrides()))
        self.assertEqual(len(d2.defensive_overrides), 2)
        self.assertNotIn("defensiveOverrides", d2.tutor_extra)

    def test_older_files_load_with_no_overrides(self):
        d = sample_deal()
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "old.json"
            save_workshop(d, p)
            data = json.loads(p.read_text(encoding="utf-8"))
            del data["defensiveOverrides"]
            p.write_text(json.dumps(data), encoding="utf-8")
            d2, _ = load_any(p)
        self.assertEqual(d2.defensive_overrides, [])
        self.assertEqual(from_tutor_dict({"south_cards": [1]}).defensive_overrides, [])

    def test_malformed_records_are_skipped(self):
        d = from_tutor_dict({"south_cards": [], "defensiveOverrides": [
            {"suitLed": "X", "occurrence": 1, "player": "East", "cardToPlay": 1},
            {"suitLed": "S", "occurrence": 1, "player": "East", "cardToPlay": "SA"},
            {"suitLed": "S", "occurrence": 1, "player": "North", "cardToPlay": 12},
            {"suitLed": "S", "occurrence": 1, "player": "East", "cardToPlay": 12}]})
        self.assertEqual(d.defensive_overrides, [DefensiveOverride("S", 1, "E", 12)])

    def test_warning_when_card_leaves_players_hand(self):
        d = self.deal_with_overrides()
        self.assertEqual(d.validation().warnings, [])
        d.swap(parse_card("DQ"), parse_card("D7"))    # West's DQ goes to East
        self.assertTrue(any(w.startswith("Override") for w in d.validation().warnings))


LIBRARY_KEYS = ["name", "description", "book", "reference", "all_cards",
                "south_cards", "west_cards", "north_cards", "east_cards",
                "contract", "contract_trump", "contract_level", "w_leads", "e_leads",
                "si_player", "si_pos", "si_suit", "si_action", "commentary",
                "page", "comment", "defensiveOverrides"]


def library_array(text):
    """The libDeals array of a library file, parsed as JSON."""
    start = text.index("const libDeals = [") + len("const libDeals = ")
    end = text.rindex("];")
    return json.loads(text[start:end + 1])


class LibraryTests(unittest.TestCase):
    EXISTING = ('//START DEAL LIST\r\nconst libDeals = [\r\n    {\r\n'
                '        "name": "Hand 4",\r\n        "comment": "uses ] and // inside"\r\n'
                '    }\r\n];\r\n//END DEAL LIST\r\n\r\nmodule.exports = { libDeals };\r\n')

    def lib_deal(self):
        d = sample_deal()
        d.author.title = "Hand 12"
        d.author.concept = "Avoidance"
        d.author.book = "Mendelson"
        d.author.page = "22"
        return d

    def test_library_dict_fields(self):
        out = to_library_dict(self.lib_deal())
        self.assertEqual(list(out), LIBRARY_KEYS)
        self.assertEqual(out["contract"], "4 Hearts")
        self.assertEqual(out["contract_trump"], "H")
        self.assertEqual(out["contract_level"], "4")
        self.assertEqual(out["w_leads"], [parse_card("DK")])
        self.assertEqual(out["si_player"], 4)
        self.assertEqual(out["all_cards"], [[], [], [], []])

    def test_append_keeps_existing_text(self):
        new = append_to_library_text(self.EXISTING, self.lib_deal())
        head = self.EXISTING[:self.EXISTING.index("    }") + 5]
        self.assertTrue(new.startswith(head + ",\r\n    {\r\n"))
        self.assertTrue(new.endswith(self.EXISTING[len(head):]))
        self.assertNotIn("\n", new.replace("\r\n", ""))     # line endings kept
        deals = library_array(new)
        self.assertEqual([d["name"] for d in deals], ["Hand 4", "Hand 12"])
        self.assertEqual(library_names(new), ["Hand 4", "Hand 12"])

    def test_append_to_empty_array_and_trailing_comma(self):
        for text in ('const libDeals = [];\n', 'const libDeals = [\n  {"name": "A"},\n];\n'):
            with self.subTest(text=text):
                new = append_to_library_text(text, self.lib_deal())
                self.assertNotIn(",,", new)
                self.assertEqual(library_array(new)[-1]["name"], "Hand 12")

    def test_rejects_file_without_lib_deals(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "other.js"
            p.write_text("const deals = [];\n", encoding="utf-8")
            with self.assertRaises(FileFormatError):
                append_to_library(self.lib_deal(), p)
            self.assertEqual(p.read_text(encoding="utf-8"), "const deals = [];\n")
        with self.assertRaises(FileFormatError):
            append_to_library_text("const libDeals = [ {", self.lib_deal())

    def test_append_to_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "crossRuffLibrary.js"
            p.write_bytes(self.EXISTING.encode("utf-8"))
            append_to_library(self.lib_deal(), p)
            append_to_library(self.lib_deal(), p)
            text = p.read_bytes().decode("utf-8")
        self.assertEqual(len(library_array(text)), 3)

    def test_real_library_file(self):
        root = Path(__file__).resolve().parents[2]
        f = root / "BWARepo" / "dealLibrary.js"
        if not f.exists():
            self.skipTest("no sample library on this machine")
        text = f.read_bytes().decode("utf-8-sig")
        new = append_to_library_text(text, self.lib_deal())
        self.assertTrue(new.startswith(text[:text.rindex("}") + 1]))
        before, after = library_array(text), library_array(new)
        self.assertEqual(after[:-1], before)
        self.assertEqual(after[-1], to_library_dict(self.lib_deal()))


class LibraryOpenTests(unittest.TestCase):
    TEXT = ('//START DEAL LIST\r\nconst libDeals = [\r\n'
            '    {\r\n        "name": "A",\r\n        "south_cards": [12]\r\n    },\r\n'
            '    // keep this comment\r\n'
            '    {\r\n        "name": "B",\r\n        "note": "has ] and , inside"\r\n    },\r\n'
            '    {\r\n        "name": "C"\r\n    }\r\n'
            '];\r\n//END DEAL LIST\r\n\r\nmodule.exports = { libDeals };\r\n')

    def names(self, text):
        return [d["name"] for d in library_deals(text)]

    def test_lists_entries(self):
        self.assertEqual(self.names(self.TEXT), ["A", "B", "C"])
        self.assertEqual(library_deals("const libDeals=[];"), [])

    def test_remove_keeps_everything_else(self):
        for index, left in [(0, ["B", "C"]), (1, ["A", "C"]), (2, ["A", "B"])]:
            with self.subTest(index=index):
                new = remove_from_library_text(self.TEXT, index)
                self.assertEqual(self.names(new), left)
                self.assertIn("// keep this comment", new)
                self.assertNotIn("\n", new.replace("\r\n", ""))   # line endings kept
                self.assertTrue(new.endswith("];\r\n//END DEAL LIST\r\n\r\n"
                                             "module.exports = { libDeals };\r\n"))
        self.assertEqual(remove_from_library_text(self.TEXT, 2),
                         self.TEXT.replace(',\r\n    {\r\n        "name": "C"\r\n    }', ""))

    def test_remove_last_with_trailing_comma_and_only_entry(self):
        new = remove_from_library_text('const libDeals = [\n  {"name": "A"},\n  {"name": "B"},\n];\n', 1)
        self.assertEqual(new, 'const libDeals = [\n  {"name": "A"},\n];\n')
        new = remove_from_library_text('const libDeals = [\n  {"name": "A"}\n];\n', 0)
        self.assertEqual(new, 'const libDeals = [\n];\n')

    def test_var_and_let_libraries_like_smallcards(self):
        for decl in ("var libDeals=[", "var libDeals = [", "let libDeals = ["):
            with self.subTest(decl=decl):
                text = self.TEXT.replace("const libDeals = [", decl)
                self.assertEqual(self.names(text), ["A", "B", "C"])
                self.assertEqual(LibraryText(text).text(), text)
                removed = remove_from_library_text(text, 1)
                self.assertIn(decl, removed)
                self.assertEqual(self.names(removed), ["A", "C"])
                appended = append_to_library_text(text, sample_deal())
                self.assertIn(decl, appended)
                self.assertEqual(len(library_deals(appended)), 4)
        with self.assertRaises(FileFormatError):
            library_deals("var otherDeals = [];")

    def test_rejects_non_json_entries(self):
        with self.assertRaises(FileFormatError):
            library_deals("const libDeals = [ {name: 'A'} ];")

    def test_library_entry_loads_like_a_tutor_file(self):
        entry = to_library_dict(sample_deal())
        deal, kind = deal_from_dict(entry)
        self.assertEqual(kind, "tutor")
        self.assertEqual(deal.hands(), sample_deal().hands())

    def test_check_out_backs_up_and_removes(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "dealLibrary.js"
            p.write_bytes(b"\xef\xbb\xbf" + self.TEXT.encode("utf-8"))
            text, bom = read_library(p)
            check_out_of_library(p, text, bom, 1)
            self.assertEqual(Path(tmp, "dealLibrary.js.bak").read_bytes(),
                             b"\xef\xbb\xbf" + self.TEXT.encode("utf-8"))
            raw = p.read_bytes()
            self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
            self.assertEqual(self.names(raw.decode("utf-8-sig")), ["A", "C"])

    def test_check_out_refuses_if_file_changed(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "dealLibrary.js"
            p.write_text(self.TEXT, encoding="utf-8", newline="")
            text, bom = read_library(p)
            changed = append_to_library_text(text, sample_deal())
            p.write_text(changed, encoding="utf-8", newline="")
            with self.assertRaises(FileFormatError):
                check_out_of_library(p, text, bom, 0)
            self.assertEqual(p.read_text(encoding="utf-8", newline=""), changed)
            self.assertFalse(Path(tmp, "dealLibrary.js.bak").exists())

    def test_real_library_files(self):
        root = Path(__file__).resolve().parents[2]
        files = [f for f in (root / "BWARepo" / "dealLibrary.js",
                             root / "InstructConstruct" / "dealLibraryBook.js") if f.exists()]
        if not files:
            self.skipTest("no sample library on this machine")
        for f in files:
            with self.subTest(file=f.name):
                text, _ = read_library(f)
                deals = library_deals(text)
                for d in deals:
                    deal_from_dict(d)
                for i in range(len(deals)):
                    self.assertEqual(library_deals(remove_from_library_text(text, i)),
                                     deals[:i] + deals[i + 1:])


class LibraryManagerTests(unittest.TestCase):
    TEXT = LibraryOpenTests.TEXT
    names = LibraryOpenTests.names

    def test_row_text_has_book_chapter_hand_only(self):
        from workshop.library_manager import manager_row_text
        entry = to_library_dict(sample_deal())      # a full entry, contract included
        entry.update(book="Mendelson", name="Hand 5")
        self.assertEqual(manager_row_text(2, entry), "2. Mendelson | Chapter ? | Hand 5")
        self.assertEqual(manager_row_text(1, {"book": "Book 1", "chapter": 2, "name": "7"}),
                         "1. Book 1 | Chapter 2 | Hand 7")
        self.assertEqual(manager_row_text(3, {"book": None}), "3. No book | Chapter ? | Hand ?")

    def test_unchanged_layout_gives_back_the_same_text(self):
        for text in (self.TEXT, "const libDeals=[];", 'const libDeals = [\n  {"name": "A"},\n];\n'):
            self.assertEqual(LibraryText(text).text(), text)

    def test_swap_keeps_formatting_and_order(self):
        lt = LibraryText(self.TEXT)
        lt.swap(0, 2)
        new = lt.text()
        self.assertEqual(self.names(new), ["C", "B", "A"])
        self.assertIn("// keep this comment", new)
        self.assertNotIn("\n", new.replace("\r\n", ""))
        self.assertTrue(new.startswith("//START DEAL LIST\r\nconst libDeals = [\r\n    {"))
        self.assertTrue(new.endswith("];\r\n//END DEAL LIST\r\n\r\nmodule.exports = { libDeals };\r\n"))
        lt.swap(0, 2)
        self.assertEqual(lt.text(), self.TEXT)

    def test_delete_matches_check_out_removal(self):
        for i in range(3):
            lt = LibraryText(self.TEXT)
            lt.delete(i)
            self.assertEqual(self.names(lt.text()), self.names(remove_from_library_text(self.TEXT, i)))
        lt = LibraryText(self.TEXT)
        for _ in range(3):
            lt.delete(0)
        self.assertEqual(library_deals(lt.text()), [])
        self.assertIn("const libDeals = [\r\n];", lt.text())

    def test_replace_library_backs_up_and_refuses_bad_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "lib.js"
            p.write_text(self.TEXT, encoding="utf-8", newline="")
            text, bom = read_library(p)
            lt = LibraryText(text)
            lt.swap(0, 1)
            with self.assertRaises(FileFormatError):     # wrong expected order: nothing written
                replace_library(p, text, bom, lt.text(), library_deals(text))
            self.assertFalse((Path(tmp) / "lib.js.bak").exists())
            deals = library_deals(text)
            replace_library(p, text, bom, lt.text(), [deals[1], deals[0], deals[2]])
            self.assertEqual(self.names(p.read_text(encoding="utf-8")), ["B", "A", "C"])
            self.assertEqual((Path(tmp) / "lib.js.bak").read_text(encoding="utf-8", newline=""),
                             self.TEXT)

    def test_real_library_files_round_trip(self):
        root = Path(__file__).resolve().parents[2]
        files = [f for f in (root / "BWARepo" / "dealLibrary.js",
                             root / "InstructConstruct" / "dealLibraryBook.js") if f.exists()]
        if not files:
            self.skipTest("no sample library on this machine")
        for f in files:
            with self.subTest(file=f.name):
                text, _ = read_library(f)
                self.assertEqual(LibraryText(text).text(), text)
                deals = library_deals(text)
                if len(deals) > 1:
                    lt = LibraryText(text)
                    lt.swap(0, len(deals) - 1)
                    self.assertEqual(library_deals(lt.text()),
                                     [deals[-1]] + deals[1:-1] + [deals[0]])


class WordExportTests(unittest.TestCase):
    def setUp(self):
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("python-docx not installed")

    def test_diagram(self):
        import docx
        from workshop.word_export import export_to_word, hand_lines
        self.assertEqual(hand_lines(from_pbn("AKJ7.Q42.KT5.")),
                         ["♠ AKJ7", "♥ Q42", "♦ K105", "♣ —"])
        d = sample_deal()
        d.contract = Contract(level=3, strain="N", declarer="N")      # East leads
        d.opening_lead = parse_card("C7")
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "deal.docx"
            export_to_word(d, p)
            doc = docx.Document(str(p))
        self.assertEqual(doc.paragraphs[0].text, "3NT by North")
        cells = [[c.text for c in row.cells] for row in doc.tables[0].rows]
        self.assertEqual(cells[0][0], "")
        self.assertEqual(cells[0][1], "♠ AQ72\n♥ K63\n♦ 984\n♣ A52")
        self.assertEqual(cells[1][2], "Opening Lead: ♣7\n\n♠ 865\n♥ 42\n♦ 765\n♣ K10987")
        self.assertEqual(cells[1][0], "♠ J94\n♥ 875\n♦ KQJ10\n♣ J63")
        self.assertEqual(cells[1][1], "N\n\nW\u2003\u2003\u2003E\n\nS")
        borders = doc.tables[0].cell(1, 1)._tc.tcPr.xml
        self.assertIn('w:val="single"', borders)


class CommandTests(unittest.TestCase):
    def test_commands(self):
        d = sample_deal()
        run_command(d, "DQ x D7")
        self.assertEqual(d.location[parse_card("DQ")], "E")
        run_command(d, "CAK P")
        self.assertIsNone(d.location[parse_card("CK")])
        run_command(d, "CK CA N")
        self.assertEqual(d.location[parse_card("CK")], "N")
        run_command(d, "S: KT3.AQJT9.A32.Q4")
        run_command(d, "lead DJ")
        self.assertEqual(d.opening_lead, parse_card("DJ"))
        with self.assertRaises(CommandError):
            run_command(d, "DQ Z")
        with self.assertRaises(CommandError):
            run_command(d, "D6 x D5")      # same hand


if __name__ == "__main__":
    unittest.main()
